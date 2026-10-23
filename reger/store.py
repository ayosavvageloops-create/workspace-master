"""Export created accounts to accounts.xlsx."""

from __future__ import annotations

from dataclasses import dataclass, fields, astuple
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from reger import config


@dataclass
class Account:
    email: str
    password: str
    phone: str
    first_name: str
    last_name: str
    proxy: str
    exit_ip: str
    profile: str
    created_at: str  # ISO-8601


_HEADERS = [f.name for f in fields(Account)]


def _ensure_file(path: Path) -> openpyxl.Workbook:
    if path.exists():
        return openpyxl.load_workbook(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accounts"
    for col, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
    return wb


def append(account: Account, path: Path | None = None) -> None:
    """Append one account row and save."""
    dest = path or config.ACCOUNTS_FILE
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = _ensure_file(dest)
    ws = wb.active
    ws.append(list(astuple(account)))
    wb.save(dest)


if __name__ == "__main__":
    import tempfile
    from datetime import datetime, timezone

    print("store self-test")

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "test_accounts.xlsx"

        acc = Account(
            email="test@example.com",
            password="secret",
            phone="+79001234567",
            first_name="Иван",
            last_name="Петров",
            proxy="socks5://proxy:1080",
            exit_ip="1.2.3.4",
            profile="profile_001",
            created_at=datetime.now(timezone.utc).isoformat(),
        )

        append(acc, path=out)
        assert out.exists()

        wb = openpyxl.load_workbook(out)
        ws = wb.active
        assert ws.max_row == 2  # header + 1 data row
        assert ws.cell(row=1, column=1).value == "email"
        assert ws.cell(row=2, column=1).value == "test@example.com"

        # Second append should create row 3
        append(acc, path=out)
        wb2 = openpyxl.load_workbook(out)
        assert wb2.active.max_row == 3

    print("  OK")
